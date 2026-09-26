"""Cover, Comps, Precedents and Summary sheets for the FPS valuation workbook."""
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.styles import Alignment
from openpyxl.utils import get_column_letter as L

from fps_data import COMPS, COMPS_ASOF, COMPS_NOTES, PRECEDENTS, PRECEDENT_NOTES


def build_rest(wb, ref, put, header_row, section, widths, st):
    MK, R, cols = ref["MK"], ref["DCF_R"], ref["DCF_COLS"]
    f_in, f_link, f_note, f_title = st["f_in"], st["f_link"], st["f_note"], st["f_title"]
    USD, USD1, USD2, NUM1, PCT, MULT = st["USD"], st["USD1"], st["USD2"], st["NUM1"], st["PCT"], st["MULT"]
    fill_key = st["fill_key"]
    fy27, fy28 = cols[0], cols[1]
    rev, ebitda = R["Revenue"], R["Adj. EBITDA"]

    # ================================================================= COMPS
    c = wb.create_sheet("Comps")
    widths(c, 30, 11, 20)
    c.column_dimensions["B"].width = 9
    put(c, "A1", "Trading Comparables - Electrical / Power Infrastructure Equipment", f_title)
    put(c, "A2", f"Market data as of {COMPS_ASOF}. $mm (local currency for non-US listings; multiples are currency-neutral). "
        "Consensus estimates are calendar-year. Blue = sourced input (see cell comments / notes below).", f_note)
    hdr = ["Company", "Ticker", "Share price", "Diluted shares (mm)", "Market cap", "Debt", "Cash", "Enterprise value",
           "LTM revenue", "LTM EBITDA", "CY26E revenue", "CY27E revenue", "CY26E EBITDA", "CY27E EBITDA",
           "EV / LTM rev", "EV / CY26E rev", "EV / LTM EBITDA", "EV / CY26E EBITDA", "EV / CY27E EBITDA",
           "CY26-27E rev growth", "CY26E EBITDA margin"]
    header_row(c, 4, hdr)
    c.row_dimensions[4].height = 42
    first = 5
    for i, co in enumerate(COMPS):
        r = first + i
        put(c, f"A{r}", co["name"])
        put(c, f"B{r}", co["ticker"])
        inputs = [("C", "price", USD2), ("D", "shares", NUM1), ("F", "debt", USD), ("G", "cash", USD),
                  ("I", "ltm_rev", USD), ("J", "ltm_ebitda", USD), ("K", "rev26", USD), ("L", "rev27", USD),
                  ("M", "ebitda26", USD), ("N", "ebitda27", USD)]
        for col, key, fmt in inputs:
            v = co.get(key)
            if v is not None:
                put(c, f"{col}{r}", v, f_in, fmt, comment=co.get("src") if col == "C" else None)
            else:
                put(c, f"{col}{r}", "n/a", f_note, align="right")
        put(c, f"E{r}", f"=C{r}*D{r}", fmt=USD)
        put(c, f"H{r}", f"=E{r}+F{r}-G{r}", fmt=USD)
        ratio = lambda num, den: f'=IF(ISNUMBER({den}{r}),IF({den}{r}>0,{num}{r}/{den}{r},"n/m"),"n/a")'
        put(c, f"O{r}", ratio("H", "I"), fmt=MULT)
        put(c, f"P{r}", ratio("H", "K"), fmt=MULT)
        put(c, f"Q{r}", ratio("H", "J"), fmt=MULT)
        put(c, f"R{r}", ratio("H", "M"), fmt=MULT)
        put(c, f"S{r}", ratio("H", "N"), fmt=MULT)
        put(c, f"T{r}", f'=IF(AND(ISNUMBER(K{r}),ISNUMBER(L{r})),L{r}/K{r}-1,"n/a")', fmt=PCT)
        put(c, f"U{r}", f'=IF(AND(ISNUMBER(K{r}),ISNUMBER(M{r})),M{r}/K{r},"n/a")', fmt=PCT)
    last = first + len(COMPS) - 1
    stats = [("Maximum", "MAX"), ("75th percentile", "Q3"), ("Mean", "AVERAGE"), ("Median", "MEDIAN"),
             ("25th percentile", "Q1"), ("Minimum", "MIN")]
    srow = last + 2
    STAT = {}
    for i, (lab, fn) in enumerate(stats):
        r = srow + i
        put(c, f"A{r}", lab, bold=True)
        for col in "OPQRSTU":
            rng = f"{col}{first}:{col}{last}"
            if fn == "Q3":
                f = f"=QUARTILE({rng},3)"
            elif fn == "Q1":
                f = f"=QUARTILE({rng},1)"
            else:
                f = f"={fn}({rng})"
            put(c, f"{col}{r}", f, fmt=PCT if col in "TU" else MULT, bold=True)
        STAT[lab] = r

    # FPS row
    fr = srow + len(stats) + 1
    put(c, f"A{fr}", "Forgent Power Solutions (at market)", bold=True)
    put(c, f"B{fr}", "FPS")
    put(c, f"C{fr}", f"={MK['Share price ($)']}", f_link, USD2)
    put(c, f"D{fr}", f"={MK['FDSO']}", f_link, NUM1)
    put(c, f"E{fr}", f"=C{fr}*D{fr}", fmt=USD)
    put(c, f"F{fr}", f"={MK['Total debt - term loan ($mm)']}", f_link, USD)
    put(c, f"G{fr}", f"={MK['Cash & equivalents ($mm)']}", f_link, USD)
    put(c, f"H{fr}", f"=E{fr}+F{fr}-G{fr}", fmt=USD)
    put(c, f"I{fr}", "=Historicals!C5", f_link, USD)
    put(c, f"J{fr}", "=Historicals!C14", f_link, USD)
    # calendarize June FYE: CY26 = 1/2 FY26A + 1/2 FY27E ; CY27 = 1/2 FY27E + 1/2 FY28E
    put(c, f"K{fr}", f"=0.5*Historicals!C5+0.5*DCF!{fy27}{rev}", fmt=USD,
        comment="Calendarized: CY2026 = 50% FY26A (Jul-25 to Jun-26) + 50% FY27E. Uses the active DCF case.")
    put(c, f"L{fr}", f"=0.5*DCF!{fy27}{rev}+0.5*DCF!{fy28}{rev}", fmt=USD, comment="CY2027 = 50% FY27E + 50% FY28E.")
    put(c, f"M{fr}", f"=0.5*Historicals!C14+0.5*DCF!{fy27}{ebitda}", fmt=USD)
    put(c, f"N{fr}", f"=0.5*DCF!{fy27}{ebitda}+0.5*DCF!{fy28}{ebitda}", fmt=USD)
    for col, num, den in [("O", "H", "I"), ("P", "H", "K"), ("Q", "H", "J"), ("R", "H", "M"), ("S", "H", "N")]:
        put(c, f"{col}{fr}", f"={num}{fr}/{den}{fr}", fmt=MULT, bold=True)
    put(c, f"T{fr}", f"=L{fr}/K{fr}-1", fmt=PCT)
    put(c, f"U{fr}", f"=M{fr}/K{fr}", fmt=PCT)

    # implied valuation
    ir = fr + 3
    section(c, ir, "Implied FPS valuation from peer multiples (25th percentile - median - 75th percentile)", 21)
    header_row(c, ir + 1, ["Metric", "", "FPS metric", "25th pct", "Median", "75th pct",
                           "EV low", "EV mid", "EV high", "Price low", "Price mid", "Price high"])
    methods = [("EV / CY26E EBITDA", "M", "R"), ("EV / CY27E EBITDA", "N", "S"), ("EV / LTM EBITDA", "J", "Q")]
    netcash = f"(-{MK['Total debt - term loan ($mm)']}+{MK['Cash & equivalents ($mm)']})"
    IMPL = {}
    for i, (lab, mcol, xcol) in enumerate(methods):
        r = ir + 2 + i
        put(c, f"A{r}", lab)
        put(c, f"C{r}", f"={mcol}{fr}", fmt=USD)
        put(c, f"D{r}", f"={xcol}{STAT['25th percentile']}", fmt=MULT)
        put(c, f"E{r}", f"={xcol}{STAT['Median']}", fmt=MULT)
        put(c, f"F{r}", f"={xcol}{STAT['75th percentile']}", fmt=MULT)
        for k, (ecol, mc) in enumerate(zip("GHI", "DEF")):
            put(c, f"{ecol}{r}", f"=$C{r}*{mc}{r}", fmt=USD)
        for pcol, ecol in zip("JKL", "GHI"):
            put(c, f"{pcol}{r}", f"=({ecol}{r}+{netcash})/{MK['FDSO']}", fmt=USD2, bold=True)
        IMPL[lab] = r
    nr = ir + 2 + len(methods) + 1
    put(c, f"A{nr}", "Read-across caveat: FPS is guiding +76% revenue growth for FY27 vs. low-double-digit growth for most peers, so "
        "peer multiples on near-year EBITDA understate a growth-adjusted value; the CY27E multiple is the more relevant "
        "anchor. Conversely FPS carries higher customer concentration (~59% data center) and a short public track record.", f_note)
    for i, n in enumerate(COMPS_NOTES):
        put(c, f"A{nr+1+i}", n, f_note)
    c.freeze_panes = "C5"

    # ============================================================ PRECEDENTS
    p = wb.create_sheet("Precedents")
    widths(p, 12, 13, 12)
    p.column_dimensions["B"].width = 20
    p.column_dimensions["C"].width = 30
    p.column_dimensions["D"].width = 44
    p.column_dimensions["J"].width = 26
    put(p, "A1", "Precedent Transactions - Electrical Distribution & Data-Center Power Infrastructure", f_title)
    put(p, "A2", "$mm. Multiples on target LTM figures unless the Basis column says otherwise. Control premia are embedded; "
        "precedent values are therefore not directly comparable to a minority trading price.", f_note)
    header_row(p, 4, ["Announced", "Acquirer", "Target", "Target description", "Enterprise value", "Target revenue",
                      "Target EBITDA", "EV / revenue", "EV / EBITDA", "Basis / notes", "In stats (1/0)",
                      "EV/rev (stats)", "EV/EBITDA (stats)"])
    p.row_dimensions[4].height = 30
    pf = 5
    for i, t in enumerate(PRECEDENTS):
        r = pf + i
        put(p, f"A{r}", t["date"])
        put(p, f"B{r}", t["acquirer"])
        put(p, f"C{r}", t["target"])
        put(p, f"D{r}", t["desc"])
        put(p, f"E{r}", t["ev"], f_in, USD, comment=t.get("src"))
        for col, key in (("F", "rev"), ("G", "ebitda")):
            if t.get(key) is not None:
                put(p, f"{col}{r}", t[key], f_in, USD)
            else:
                put(p, f"{col}{r}", "n/a", f_note, align="right")
        # disclosed multiples without disclosed base figures are entered directly
        if t.get("ev_rev") is not None:
            put(p, f"H{r}", t["ev_rev"], f_in, MULT)
        else:
            put(p, f"H{r}", f'=IF(ISNUMBER(F{r}),E{r}/F{r},"n/a")', fmt=MULT)
        if t.get("ev_ebitda") is not None:
            put(p, f"I{r}", t["ev_ebitda"], f_in, MULT)
        else:
            put(p, f"I{r}", f'=IF(ISNUMBER(G{r}),E{r}/G{r},"n/a")', fmt=MULT)
        put(p, f"J{r}", t.get("basis", "LTM"))
        put(p, f"K{r}", t["include"], f_in, "0", align="center")
        put(p, f"L{r}", f'=IF(AND(K{r}=1,ISNUMBER(H{r})),H{r},"")', fmt=MULT)
        put(p, f"M{r}", f'=IF(AND(K{r}=1,ISNUMBER(I{r})),I{r},"")', fmt=MULT)
        for col in "BCDJ":
            p[f"{col}{r}"].alignment = Alignment(wrap_text=True, vertical="top")
    pl = pf + len(PRECEDENTS) - 1
    ps = pl + 2
    PST = {}
    for i, (lab, fn) in enumerate(stats):
        r = ps + i
        put(p, f"C{r}", lab, bold=True)
        for col, scol in (("H", "L"), ("I", "M")):
            rng = f"{scol}{pf}:{scol}{pl}"
            f = {"Q3": f"=QUARTILE({rng},3)", "Q1": f"=QUARTILE({rng},1)"}.get(fn, f"={fn}({rng})")
            put(p, f"{col}{r}", f, fmt=MULT, bold=True)
        PST[lab] = r
    pr = ps + len(stats) + 1
    section(p, pr, "Implied FPS valuation (applied to FPS CY26E, 25th pct - median - 75th pct; core deals only)", 13)
    header_row(p, pr + 1, ["Metric", "", "FPS CY26E metric", "25th pct", "Median", "75th pct", "Price low", "Price mid", "Price high"])
    PIMPL = {}
    for i, (lab, mref, xcol) in enumerate([("EV / EBITDA", f"=Comps!M{fr}", "I"),
                                           ("EV / revenue", f"=Comps!K{fr}", "H")]):
        r = pr + 2 + i
        put(p, f"A{r}", lab)
        put(p, f"C{r}", mref, f_link, USD, comment="Calendar-2026E (50% FY26A + 50% FY27E, active DCF case).")
        put(p, f"D{r}", f"={xcol}{PST['25th percentile']}", fmt=MULT)
        put(p, f"E{r}", f"={xcol}{PST['Median']}", fmt=MULT)
        put(p, f"F{r}", f"={xcol}{PST['75th percentile']}", fmt=MULT)
        for pcol, mc in zip("GHI", "DEF"):
            put(p, f"{pcol}{r}", f"=($C{r}*{mc}{r}+{netcash})/{MK['FDSO']}", fmt=USD2, bold=True)
        PIMPL[lab] = r
    for i, n in enumerate(PRECEDENT_NOTES):
        put(p, f"A{pr+5+i}", n, f_note)
    p.freeze_panes = "D5"

    # =============================================================== SUMMARY
    sm = wb.create_sheet("Summary")
    widths(sm, 44, 13, 8)
    put(sm, "A1", "Valuation Summary - Football Field ($ per FPS share)", f_title)
    put(sm, "A2", '="Active DCF case: "&' + ref["SCNAME"] + '&"   |   Current price $"&TEXT(' + MK["Share price ($)"] + ',"0.00")', f_note)
    header_row(sm, 4, ["Methodology", "Low", "High", "Spread", "Midpoint", "vs. current"])
    sens = "Sensitivity!"
    rows = [
        ("DCF - perpetuity growth (WACC +/-1%, g 2.5-3.5%)", f"=MIN({sens}D7:F9)", f"=MAX({sens}D7:F9)"),
        ("DCF - exit multiple (WACC +/-1%, 12-16x)", f"=MIN({sens}D15:F17)", f"=MAX({sens}D15:F17)"),
        ("Trading comps - EV / CY27E EBITDA (25th-75th)", f"=Comps!J{IMPL['EV / CY27E EBITDA']}", f"=Comps!L{IMPL['EV / CY27E EBITDA']}"),
        ("Trading comps - EV / CY26E EBITDA (25th-75th)", f"=Comps!J{IMPL['EV / CY26E EBITDA']}", f"=Comps!L{IMPL['EV / CY26E EBITDA']}"),
        ("Precedent transactions - EV / CY26E EBITDA (25th-75th)", f"=Precedents!G{PIMPL['EV / EBITDA']}", f"=Precedents!I{PIMPL['EV / EBITDA']}"),
        ("Analyst price targets (consensus mean)", 57.90, 57.90),
    ]
    for i, (lab, lo, hi) in enumerate(rows):
        r = 5 + i
        put(sm, f"A{r}", lab)
        is_in = not isinstance(lo, str)
        put(sm, f"B{r}", lo, f_in if is_in else f_link, USD2,
            comment="Consensus mean target $57.90 (stockanalysis.com, 25-Sep-2026; 10 analysts). Reference only." if is_in else None)
        put(sm, f"C{r}", hi, f_in if is_in else f_link, USD2)
        put(sm, f"D{r}", f"=C{r}-B{r}", fmt=USD2)
        put(sm, f"E{r}", f"=(B{r}+C{r})/2", fmt=USD2)
        put(sm, f"F{r}", f"=E{r}/{MK['Share price ($)']}-1", fmt=PCT)
    n = len(rows)
    put(sm, f"A{6+n}", "Current share price", bold=True)
    put(sm, f"B{6+n}", f"={MK['Share price ($)']}", f_link, USD2, fill=fill_key)
    put(sm, f"A{7+n}", "DCF point estimate - perpetuity / exit multiple", bold=True)
    put(sm, f"B{7+n}", f"={ref['DCF_PRICE_PG']}", f_link, USD2)
    put(sm, f"C{7+n}", f"={ref['DCF_PRICE_EX']}", f_link, USD2)
    put(sm, f"A{8+n}", "Market-implied EV / FY27E EBITDA (guidance midpoint)", bold=True)
    put(sm, f"B{8+n}", f"={MK['EV']}/DCF!{fy27}{ebitda}", fmt=MULT)

    ch = BarChart()
    ch.type = "bar"
    ch.grouping = "stacked"
    ch.overlap = 100
    ch.title = "FPS valuation range ($/share)"
    ch.y_axis.title = "$ per share"
    data = Reference(sm, min_col=2, min_row=4, max_row=4 + n)
    spread = Reference(sm, min_col=4, min_row=4, max_row=4 + n)
    cats = Reference(sm, min_col=1, min_row=5, max_row=4 + n)
    ch.add_data(data, titles_from_data=True)
    ch.add_data(spread, titles_from_data=True)
    ch.set_categories(cats)
    ch.series[0].graphicalProperties.noFill = True
    ch.series[0].graphicalProperties.line.noFill = True
    ch.series[1].graphicalProperties.solidFill = "1F3864"
    ch.legend = None
    ch.height, ch.width = 9, 22
    ch.x_axis.scaling.orientation = "maxMin"
    ch.x_axis.delete = False
    ch.y_axis.delete = False
    sm.add_chart(ch, f"A{11+n}")

    # ================================================================= COVER
    cv = wb.create_sheet("Cover", 0)
    widths(cv, 110, 10, 2)
    put(cv, "A1", "Forgent Power Solutions, Inc. (NYSE: FPS) - Valuation Model", f_title)
    lines = [
        "DCF (10-year, FY27E-FY36E, Base/Bull/Bear), trading comparables, precedent transactions, football-field summary.",
        "",
        "SHEETS",
        "  Inputs - market data, capital structure, scenario drivers (scenario switch in yellow). Edit blue cells only.",
        "  Historicals - FY25A/FY26A and FY26 quarters from the FY2026 10-K and Q4 earnings release.",
        "  WACC - CAPM cost of equity with peer-based beta; optional override cell.",
        "  DCF - unlevered FCF build, perpetuity-growth and exit-multiple terminal values, equity bridge.",
        "  Sensitivity - share price vs WACC x terminal growth and WACC x exit multiple (live formulas).",
        "  Comps - peer trading multiples with 25th/median/75th percentile implied FPS values.",
        "  Precedents - M&A multiples for power-distribution / data-center power targets.",
        "  Summary - football field across methods.",
        "",
        "KEY MODELLING CHOICES",
        "  * FY27E anchored to company guidance (revenue $2.4-2.6bn, adj. EBITDA $575-625mm). FY28E onward are analyst assumptions, not guidance.",
        "  * SBC is deducted as a cash-equivalent cost even though company 'Adjusted EBITDA' adds it back.",
        "  * Up-C: Class B units are exchangeable 1:1, so they are counted in diluted shares and no NCI is deducted.",
        "  * Tax receivable agreement (TRA): cash flows are taxed at a full 25% rate with no benefit from the IPO basis step-up. "
        "Since FPS pays 85% of that step-up benefit to legacy holders, ignoring both the shield and the TRA liability is roughly neutral "
        "(it slightly understates value by the retained 15%).",
        "  * Operating leases are left inside EBITDA and excluded from net debt (consistent treatment).",
        "  * Acquired-intangible amortization is not tax-effected (conservative).",
        "",
        "COLOUR CODE: blue = hardcoded input, black = formula, green = link to another sheet, yellow fill = key lever / output.",
        "Sources: SEC EDGAR filings for CIK 0002080126 (10-K FY2026, 8-K earnings releases, 10-Qs, S-1/424B4); market data per cell comments.",
        "Not investment advice. Peer and deal data are compiled from public sources and should be re-verified before use.",
    ]
    for i, t in enumerate(lines):
        put(cv, f"A{3+i}", t, st["f_calc"], bold=t in ("SHEETS", "KEY MODELLING CHOICES"))
        cv[f"A{3+i}"].alignment = Alignment(wrap_text=True)
