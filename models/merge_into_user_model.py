"""Merge forward comps, precedent transactions and a football field into the user's FPS model.

Input : Forgent_DCF_F26_iPhone.xlsx (user's 3-statement DCF)
Output: Forgent_DCF_F26_merged.xlsx
Changes to existing sheets are limited to: Assumptions!D40 (exit-multiple override) and its note,
Assumptions!L9:L11 (re-pasted scenario DCF prices, set by a later step), and Notes text.
"""
import sys
from copy import copy

from openpyxl import load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import get_column_letter as L

sys.path.insert(0, __file__.rsplit("/", 1)[0])
from fps_data import COMPS, PRECEDENTS  # noqa: E402

SRC, DST = sys.argv[1], sys.argv[2]
EXIT_MULT = 14.0

wb = load_workbook(SRC)

# ---------- house style (matches the user's workbook) ----------
FN = "Garamond"
f_title = Font(name=FN, size=14, bold=True, color="FF6F0505")
f_norm = Font(name=FN, size=11)
f_bold = Font(name=FN, size=11, bold=True)
f_in = Font(name=FN, size=11, color="FF0000FF")
f_link = Font(name=FN, size=11, color="FF008000")
f_note = Font(name=FN, size=9, color="FF7F7F7F")
hdr_border = Border(bottom=Side(style="thin"))
USD2 = '"$"#,##0.00_);\\("$"#,##0.00\\)'
USD0 = '#,##0_);\\(#,##0\\);"-"'
MULT = '0.0\\x'
PCT = '0.0%;\\(0.0%\\);"-"'


def put(ws, ref, v, font=f_norm, fmt=None, comment=None, wrap=False, align=None):
    c = ws[ref]
    c.value = v
    c.font = copy(font)
    if fmt:
        c.number_format = fmt
    if comment:
        c.comment = Comment(comment, "KJ")
    if wrap or align:
        c.alignment = Alignment(wrap_text=wrap, horizontal=align, vertical="top" if wrap else None)
    return c


def headers(ws, row, labels, start=2):
    for i, lab in enumerate(labels):
        c = ws.cell(row=row, column=start + i, value=lab)
        c.font = copy(f_bold)
        c.border = hdr_border
        c.alignment = Alignment(horizontal="center" if i else "left", wrap_text=True, vertical="bottom")


def new_sheet(title, before="Notes"):
    ws = wb.create_sheet(title, wb.sheetnames.index(before))
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 2.96
    return ws


STATS = [("Max", "MAX"), ("75th percentile", "Q3"), ("Median", "MEDIAN"), ("Mean", "AVERAGE"),
         ("25th percentile", "Q1"), ("Min", "MIN")]


def stat_formula(fn, rng):
    return {"Q3": f"=QUARTILE({rng},3)", "Q1": f"=QUARTILE({rng},1)"}.get(fn, f"={fn}({rng})")


# =================================================================== 1. exit multiple
a = wb["Assumptions"]
put(a, "D40", EXIT_MULT, f_in, a["D40"].number_format)
a["D40"].fill = copy(a["D40"].fill)
put(a, "H40", "14.0x = mature-company exit for FY36 (growth ~3%): between precedent median (~12.6x) and "
    "25th-pct peer EV / CY27E EBITDA (~15.5x) on 'Fwd Comps'. Blank reverts to peer EV/LTM median (27.8x), "
    "which implied ~7% perpetual growth (DCF!E47).", f_note)

# =================================================================== 2. Fwd Comps
fc = new_sheet("Fwd Comps")
put(fc, "A1", '="Forward Trading Comparables for "&Name', f_title)
put(fc, "A2", "(Dollars in Millions except share price; SU, LR in EUR and HPS.A in CAD - multiples are currency-neutral)", f_norm)
put(fc, "B3", "Adds consensus CY26E/CY27E revenue and EBITDA (MarketScreener), which 'Comps Analysis' lacks. "
    "Prices and balance sheets as of 9/25/2026 close (StockAnalysis).", f_note)
cols = ["Company", "Ticker", "Price", "Shares (MM)", "Mkt cap", "Total debt", "Cash", "Enterprise value",
        "CY26E revenue", "CY27E revenue", "CY26E EBITDA", "CY27E EBITDA", "EV / CY26E revenue",
        "EV / CY26E EBITDA", "EV / CY27E EBITDA", "CY26-27E revenue growth", "CY26E EBITDA margin", "Notes"]
headers(fc, 5, cols)
fc.row_dimensions[5].height = 45
widths = [26, 9, 10, 10, 11, 10, 10, 11, 10, 10, 10, 10, 10, 10, 10, 10, 10, 60]
for i, w in enumerate(widths):
    fc.column_dimensions[L(2 + i)].width = w
first = 6
for i, co in enumerate(COMPS):
    r = first + i
    put(fc, f"B{r}", co["name"])
    put(fc, f"C{r}", co["ticker"])
    for col, key, fmt in [("D", "price", USD2), ("E", "shares", "#,##0.00"), ("G", "debt", USD0), ("H", "cash", USD0),
                          ("J", "rev26", USD0), ("K", "rev27", USD0), ("L", "ebitda26", USD0), ("M", "ebitda27", USD0)]:
        put(fc, f"{col}{r}", co[key], f_in, fmt)
    put(fc, f"F{r}", f"=D{r}*E{r}", fmt=USD0)
    put(fc, f"I{r}", f"=F{r}+G{r}-H{r}", fmt=USD0)
    put(fc, f"N{r}", f"=I{r}/J{r}", fmt=MULT)
    put(fc, f"O{r}", f"=I{r}/L{r}", fmt=MULT)
    put(fc, f"P{r}", f"=I{r}/M{r}", fmt=MULT)
    put(fc, f"Q{r}", f"=K{r}/J{r}-1", fmt=PCT)
    put(fc, f"R{r}", f"=L{r}/J{r}", fmt=PCT)
    put(fc, f"S{r}", co["src"], f_note)
last = first + len(COMPS) - 1
srow = last + 2
FST = {}
for i, (lab, fn) in enumerate(STATS):
    r = srow + i
    put(fc, f"B{r}", lab, f_bold)
    for col in "NOPQR":
        put(fc, f"{col}{r}", stat_formula(fn, f"{col}{first}:{col}{last}"), f_bold, PCT if col in "QR" else MULT)
    FST[lab] = r
fr = srow + len(STATS) + 1
put(fc, f"B{fr}", "=Name", f_bold)
put(fc, f"C{fr}", "=Ticker")
put(fc, f"D{fr}", "=shareprice", fmt=USD2)
put(fc, f"E{fr}", "=DilutedShares", fmt="#,##0.00")
put(fc, f"F{fr}", "=marketcap", fmt=USD0)
put(fc, f"G{fr}", "=totaldebt", fmt=USD0)
put(fc, f"H{fr}", "=cash", fmt=USD0)
put(fc, f"I{fr}", "=Enterprise", fmt=USD0, comment="Named range Enterprise: includes net TRA claim, consistent with 'Comps Analysis'.")
put(fc, f"J{fr}", "=0.5*Inc_Stmt!E12+0.5*Inc_Stmt!F12", fmt=USD0,
    comment="Calendarized from June FYE: CY26 = 50% FY26A + 50% FY27E. Follows the live scenario.")
put(fc, f"K{fr}", "=0.5*Inc_Stmt!F12+0.5*Inc_Stmt!G12", fmt=USD0, comment="CY27 = 50% FY27E + 50% FY28E.")
put(fc, f"L{fr}", "=0.5*Inc_Stmt!E43+0.5*Inc_Stmt!F43", fmt=USD0,
    comment="Company-definition adj. EBITDA (adds back SBC and one-offs), comparable to peers' consensus adjusted EBITDA.")
put(fc, f"M{fr}", "=0.5*Inc_Stmt!F43+0.5*Inc_Stmt!G43", fmt=USD0)
for col, num, den in [("N", "I", "J"), ("O", "I", "L"), ("P", "I", "M")]:
    put(fc, f"{col}{fr}", f"={num}{fr}/{den}{fr}", f_bold, MULT)
put(fc, f"Q{fr}", f"=K{fr}/J{fr}-1", fmt=PCT)
put(fc, f"R{fr}", f"=L{fr}/J{fr}", fmt=PCT)

ir = fr + 3
put(fc, f"B{ir}", '="Implied valuation of "&Name', f_bold)
headers(fc, ir + 1, ["Multiple", "", "FPS metric", "25th pct", "Median", "75th pct", "Price @ 25th", "Price @ median", "Price @ 75th"])
FIMPL = {}
for i, (lab, mcol, xcol) in enumerate([("EV / CY26E EBITDA", "L", "O"), ("EV / CY27E EBITDA", "M", "P"),
                                       ("EV / CY26E revenue", "J", "N")]):
    r = ir + 2 + i
    put(fc, f"B{r}", lab)
    put(fc, f"D{r}", f"={mcol}{fr}", fmt=USD0)
    for col, st in zip("EFG", ("25th percentile", "Median", "75th percentile")):
        put(fc, f"{col}{r}", f"={xcol}{FST[st]}", fmt=MULT)
    for pcol, mc in zip("HIJ", "EFG"):
        put(fc, f"{pcol}{r}", f"=({mc}{r}*$D{r}-NetClaims)/DilutedShares", f_bold, USD2)
    FIMPL[lab] = r
nr = ir + 6
notes = [
    "EV-based prices subtract NetClaims (debt - cash + TRA - DTA - restricted cash), the same bridge as 'Comps Analysis'.",
    "CY27E is the more relevant anchor: FPS guides +76% FY27 revenue growth vs. ~10-30% for peers, so CY26E multiples understate a growth-adjusted value.",
    "LTM/consensus caveats: GE Vernova EBITDA depressed by Wind; Hammond is pro forma for AEG (closed 6/29/2026, ~C$410m acquisition debt added); "
    "Powell FY ends Sep (not calendarized). MarketScreener consensus date could not be confirmed; CY26 revenue matched StockAnalysis at 9/25/2026.",
]
for i, t in enumerate(notes):
    put(fc, f"B{nr+i}", t, f_note)
fc.freeze_panes = "D6"

# =================================================================== 3. Precedents
pr = new_sheet("Precedents")
put(pr, "A1", '="Precedent Transactions for "&Name', f_title)
put(pr, "A2", "(Dollars in Millions; multiples on the basis each acquirer disclosed - see Basis column)", f_norm)
put(pr, "B3", "Includes a control premium, so values are not directly comparable to a minority trading price. "
    "'In stats' = 1 for electrical distribution / modular power / grid equipment; adjacent deals shown for context only.", f_note)
pcols = ["Announced", "Acquirer", "Target", "Description", "Enterprise value", "Target revenue", "Target EBITDA",
         "EV / revenue", "EV / EBITDA", "Basis", "In stats", "EV/rev (stats)", "EV/EBITDA (stats)", "Source"]
headers(pr, 5, pcols)
pr.row_dimensions[5].height = 32
for i, w in enumerate([10, 14, 24, 40, 10, 10, 10, 9, 9, 42, 7, 9, 9, 40]):
    pr.column_dimensions[L(2 + i)].width = w
pf = 6
for i, t in enumerate(PRECEDENTS):
    r = pf + i
    put(pr, f"B{r}", t["date"])
    put(pr, f"C{r}", t["acquirer"], wrap=True)
    put(pr, f"D{r}", t["target"], wrap=True)
    put(pr, f"E{r}", t["desc"], wrap=True)
    put(pr, f"F{r}", t["ev"], f_in, USD0)
    for col, key in (("G", "rev"), ("H", "ebitda")):
        if t.get(key) is not None:
            put(pr, f"{col}{r}", t[key], f_in, USD0)
        else:
            put(pr, f"{col}{r}", "N/A", align="right")
    put(pr, f"I{r}", t["ev_rev"] if t.get("ev_rev") else f'=IF(ISNUMBER(G{r}),F{r}/G{r},"N/A")',
        f_in if t.get("ev_rev") else f_norm, MULT)
    put(pr, f"J{r}", t["ev_ebitda"] if t.get("ev_ebitda") else f'=IF(ISNUMBER(H{r}),F{r}/H{r},"N/A")',
        f_in if t.get("ev_ebitda") else f_norm, MULT)
    put(pr, f"K{r}", t["basis"], f_note, wrap=True)
    put(pr, f"L{r}", t["include"], f_in, "0", align="center")
    put(pr, f"M{r}", f'=IF(AND(L{r}=1,ISNUMBER(I{r})),I{r},"")', fmt=MULT)
    put(pr, f"N{r}", f'=IF(AND(L{r}=1,ISNUMBER(J{r})),J{r},"")', fmt=MULT)
    put(pr, f"O{r}", t["src"], f_note, wrap=True)
pl = pf + len(PRECEDENTS) - 1
ps = pl + 2
PST = {}
for i, (lab, fn) in enumerate(STATS):
    r = ps + i
    put(pr, f"D{r}", lab, f_bold)
    put(pr, f"I{r}", stat_formula(fn, f"M{pf}:M{pl}"), f_bold, MULT)
    put(pr, f"J{r}", stat_formula(fn, f"N{pf}:N{pl}"), f_bold, MULT)
    PST[lab] = r
pi = ps + len(STATS) + 1
put(pr, f"B{pi}", '="Implied valuation of "&Name&" (core deals)"', f_bold)
headers(pr, pi + 1, ["Multiple", "", "FPS CY26E metric", "25th pct", "Median", "75th pct", "Price @ 25th", "Price @ median", "Price @ 75th"])
PIMPL = {}
for i, (lab, mref, xcol) in enumerate([("EV / EBITDA", f"='Fwd Comps'!L{fr}", "J"), ("EV / revenue", f"='Fwd Comps'!J{fr}", "I")]):
    r = pi + 2 + i
    put(pr, f"B{r}", lab)
    put(pr, f"D{r}", mref, f_link, USD0,
        comment="Most deals are priced on current/next-year EBITDA, so they are applied to FPS calendar-2026E, not FY26A LTM.")
    for col, st in zip("EFG", ("25th percentile", "Median", "75th percentile")):
        put(pr, f"{col}{r}", f"={xcol}{PST[st]}", fmt=MULT)
    for pcol, mc in zip("HIJ", "EFG"):
        put(pr, f"{pcol}{r}", f"=({mc}{r}*$D{r}-NetClaims)/DilutedShares", f_bold, USD2)
    PIMPL[lab] = r
put(pr, f"B{pi+5}", "Forgent's own roll-up deals under Neos (MGM Transformer $424.7m, VanTran $432.7m, PwrQ $103.0m, "
    "States Manufacturing $68.5m; 2023-24) disclose price only, so no multiples (424B4).", f_note)
put(pr, f"B{pi+6}", "Revenue multiples mix LTM and forward bases and span 2.6-6.3x; treat the EV/revenue row as a cross-check only.", f_note)
pr.freeze_panes = "E6"

# =================================================================== 4. Football Field
ff = new_sheet("Football Field")
put(ff, "A1", '="Valuation Summary for "&Name', f_title)
put(ff, "A2", '="Value per share; current price $"&TEXT(shareprice,"0.00")&"; live DCF scenario: "&Assumptions!D45', f_norm)
headers(ff, 4, ["Methodology", "Low", "High", "Spread", "Midpoint", "Midpoint vs. current"])
ff.column_dimensions["B"].width = 58
for c in "CDEFG":
    ff.column_dimensions[c].width = 12
rows = [
    ("DCF - Gordon growth only (WACC +/-0.5%, g 2.5-3.5%)", "=MIN(DCF!G56:I58)", "=MAX(DCF!G56:I58)"),
    ("DCF - exit multiple only (WACC +/-0.5%, exit +/-3x)", "=MIN(DCF!G65:I67)", "=MAX(DCF!G65:I67)"),
    ("DCF - blended TV, Bear to Bull scenario", "=MIN(Assumptions!L9:L11)", "=MAX(Assumptions!L9:L11)"),
    ("Trading comps - NTM P/E (25th-75th)", "='Comps Analysis'!G44", "='Comps Analysis'!G42"),
    ("Forward comps - EV / CY27E EBITDA (25th-75th)", f"='Fwd Comps'!H{FIMPL['EV / CY27E EBITDA']}", f"='Fwd Comps'!J{FIMPL['EV / CY27E EBITDA']}"),
    ("Forward comps - EV / CY26E EBITDA (25th-75th)", f"='Fwd Comps'!H{FIMPL['EV / CY26E EBITDA']}", f"='Fwd Comps'!J{FIMPL['EV / CY26E EBITDA']}"),
    ("Precedent transactions - EV / CY26E EBITDA (25th-75th)", f"=Precedents!H{PIMPL['EV / EBITDA']}", f"=Precedents!J{PIMPL['EV / EBITDA']}"),
    ("Since-IPO trading range", "=low", "=high"),
    ("Analyst consensus price target (reference)", 57.90, 57.90),
]
for i, (lab, lo, hi) in enumerate(rows):
    r = 5 + i
    put(ff, f"B{r}", lab)
    is_in = not isinstance(lo, str)
    cm = "Mean 12-month target, 10 analysts; stockanalysis.com, 9/25/2026." if is_in else None
    put(ff, f"C{r}", lo, f_in if is_in else f_link, USD2, comment=cm)
    put(ff, f"D{r}", hi, f_in if is_in else f_link, USD2)
    put(ff, f"E{r}", f"=D{r}-C{r}", fmt=USD2)
    put(ff, f"F{r}", f"=AVERAGE(C{r},D{r})", fmt=USD2)
    put(ff, f"G{r}", f"=F{r}/shareprice-1", fmt=PCT)
n = len(rows)
put(ff, f"B{6+n}", "Current share price", f_bold)
put(ff, f"C{6+n}", "=shareprice", f_bold, USD2)
put(ff, f"B{7+n}", "Target price (base DCF, blended TV)", f_bold)
put(ff, f"C{7+n}", "=TargetPrice", f_bold, USD2)
put(ff, f"B{8+n}", "Market-implied EV / FY27E adj. EBITDA", f_bold)
put(ff, f"C{8+n}", "=Enterprise/Inc_Stmt!F43", f_bold, MULT)
put(ff, f"B{9+n}", "DCF rows follow the live scenario except the Bear-Bull row, which uses the pasted values in Assumptions!L9:L11.", f_note)

ch = BarChart()
ch.type = "bar"
ch.grouping = "stacked"
ch.overlap = 100
ch.title = "Forgent (FPS) - value per share"
ch.add_data(Reference(ff, min_col=3, min_row=4, max_row=4 + n), titles_from_data=True)
ch.add_data(Reference(ff, min_col=5, min_row=4, max_row=4 + n), titles_from_data=True)
ch.set_categories(Reference(ff, min_col=2, min_row=5, max_row=4 + n))
ch.series[0].graphicalProperties.noFill = True
ch.series[0].graphicalProperties.line.noFill = True
ch.series[1].graphicalProperties.solidFill = "6F0505"
ch.legend = None
ch.x_axis.scaling.orientation = "maxMin"
ch.x_axis.delete = False
ch.y_axis.delete = False
ch.y_axis.numFmt = '"$"0'
ch.height, ch.width = 9, 22
ff.add_chart(ch, f"B{11+n}")

# =================================================================== 5. Notes
nt = wb["Notes"]
put(nt, "B17", "Exit multiple: overridden to 14.0x at Assumptions!D40 (was the peer median EV / LTM EBITDA of 27.8x, "
    "a near-peak AI-cycle trailing multiple that implied ~7% perpetual growth when applied to FY36). 14.0x sits between the "
    "precedent median (~12.6x) and the 25th-pct peer EV / CY27E EBITDA (~15.5x). Check DCF!E47 after any change.",
    copy(nt["B17"].font))
base_font = copy(nt["B19"].font)
put(nt, "B24", "Merged additions (Sept 2026)", copy(nt["B18"].font))
adds = [
    "'Fwd Comps': 9 peers with consensus CY26E/CY27E revenue and EBITDA (MarketScreener); FPS calendarized from June FYE.",
    "'Precedents': 11 core power-distribution / data-center power deals (2021-2026) + 4 adjacent deals excluded from stats; sources per row.",
    "'Football Field': DCF, trading comps, forward comps, precedents, trading range and analyst target on one page.",
    "Assumptions!L9:L11 re-pasted after the exit-multiple change (Base / Bull / Bear).",
]
for i, t in enumerate(adds):
    put(nt, f"B{25+i}", t, base_font)

wb.save(DST)
print("saved", DST)
